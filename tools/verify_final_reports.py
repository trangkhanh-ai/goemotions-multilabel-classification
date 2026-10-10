"""Kiểm các báo cáo đã xuất thật, lưu checksum và tạo ảnh trang để đọc lại.

Không chạy Word, mô hình, GPU hoặc test huấn luyện. Ảnh cần được xem bằng mắt;
script chỉ ghi pending_manual_review, không tự tuyên bố đã kiểm hình thức.
"""

import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
import pypdfium2 as pdfium


ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
PREVIEWS = REPORTS / "execution/final_report_preview"
NAMES = ("BAO_CAO_DO_AN_GOEMOTIONS_IEEE", "BAI_BAO_GOEMOTIONS_IEEE", "BAO_CAO_TIEN_DO_1", "BAO_CAO_TIEN_DO_2")
SOURCES = ("BAO_CAO_DO_AN_NOI_DUNG.md", "BAI_BAO_GOEMOTIONS_IEEE_NOI_DUNG.md", "BAO_CAO_TIEN_DO_1.md", "BAO_CAO_TIEN_DO_2.md")
MEMBERS = ("Bảo Duy Nguyễn", "Quốc Khánh", "Đức Trí", "Nhật Huy")


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def compact(text):
    return re.sub(r"\s+", "", text)


def citation_order(source):
    body = source.split("# TÀI LIỆU THAM KHẢO", 1)[0].split("## Tài liệu tham khảo", 1)[0]
    return list(dict.fromkeys(int(x) for x in re.findall(r"\[(\d{1,2})(?:\]|,)", body) if int(x) > 0))


def preview(document, index, label, stem):
    path = PREVIEWS / f"{stem}_{label}_p{index+1}.png"
    page = document[index]
    bitmap = page.render(scale=1.25)
    image = bitmap.to_pil()
    image.save(path)
    image.close()
    bitmap.close()
    page.close()
    return {"page": index + 1, "label": label, "path": str(path.relative_to(ROOT)), "sha256": digest(path)}


def inspect(name, source_name, expected_refs):
    docx_path, pdf_path, source_path = REPORTS / f"{name}.docx", REPORTS / f"{name}.pdf", REPORTS / source_name
    source = source_path.read_text(encoding="utf-8")
    word = Document(docx_path)
    cols = []
    for section in word.sections:
        node = section._sectPr.find(qn("w:cols"))
        cols.append(int(node.get(qn("w:num"), "1")) if node is not None else 1)
    pdf = pdfium.PdfDocument(pdf_path)
    page_texts, sizes = [], []
    for index in range(len(pdf)):
        page = pdf[index]
        sizes.append(list(page.get_size()))
        text_page = page.get_textpage()
        page_texts.append(text_page.get_text_range())
        text_page.close()
        page.close()
    full = "\n".join(page_texts)
    full_compact = compact(full)
    checks = {
        "pdf_has_text_every_page": all(len(compact(text)) > 20 for text in page_texts),
        "letter_page_size": all(abs(width - 612) < 1 and abs(height - 792) < 1 for width, height in sizes),
        "pdf_exported_after_docx_save": pdf_path.stat().st_mtime >= docx_path.stat().st_mtime,
        "four_known_members": all(compact(member) in full_compact for member in MEMBERS),
        "no_removed_member_or_template_lecturer": not any(x in full for x in ("Hoàng Phúc", "Huỳnh Thành Lộc")),
        "no_word_field_error": not any(x in full for x in ("Error! Reference", "Error! Bookmark", "Lỗi! Không tìm thấy")),
        "citation_first_use_order": citation_order(source) == list(range(1, expected_refs + 1)),
        "current_complete_c_runs": "9/9" in source or "ba seed 42, 123, 2026" in source,
    }
    if name.startswith("BAO_CAO_TIEN"):
        checks["blank_admin_fields"] = "____________________" in source
        checks["not_claimed_submitted"] = "không xác nhận đã nộp" in source
    elif name.startswith("BAO_CAO_DO_AN"):
        checks.update(
            six_chapters=all(f"CHƯƠNG {index}." in source for index in range(1, 7)),
            blank_admin_fields=all(f"**{label}:** ____________________" in source for label in
                                   ("Giảng viên hướng dẫn", "Mã lớp học phần", "Năm học / học kỳ", "Mã số sinh viên")),
            real_demo_figure="Hình5.4." in full_compact,
            real_case_ids=all(sample_id in full for sample_id in ("eczj48j", "ed0jr9i", "eczcvgx")),
            rare_negative_kept="-0.0257" in full,
            selected_test_mean_std=all(value in full_compact for value in ("0.4720±0.0045", "0.5038±0.0097")),
            error_group_counts_reference="reports/errors_test_standard_fixed/counts.csv" in full_compact
                                         and "group_summary.csv" not in full,
        )
    else:
        checks.update(
            body_two_columns=2 in cols,
            original_foundation_not_recent_substitute="2020" in source,
            real_case_ids=all(sample_id in full for sample_id in ("eczj48j", "ed0jr9i", "eczcvgx")),
            rare_negative_kept="-0.0257" in full,
            selected_test_mean_std=all(value in full_compact for value in ("0.4720±0.0045", "0.5038±0.0097")),
            micro_tradeoff_kept="0.6048" in full and "0.5900" in full,
        )
    if name in ("BAO_CAO_DO_AN_GOEMOTIONS_IEEE", "BAI_BAO_GOEMOTIONS_IEEE"):
        checks.update(
            three_application_scenarios=all(compact(term) in full_compact for term in
                ("định tuyến", "khủng hoảng thương hiệu", "rà soát bình luận", "MTTR")),
            hyperparameter_search_not_claimed=compact("chưa thực hiện tìm kiếm siêu tham số có hệ thống") in full_compact,
            validation_test_roles_explicit=all(compact(term) in full_compact for term in
                ("validation tuned F1", "test locked F1")),
            observed_bert_precision_recall_hamming=all(value in full for value in
                ("0.5322", "0.6437", "0.6421", "0.5446", "0.0318", "0.0373")),
            balanced_counterexample_kept=all(value in full for value in
                ("0.4043", "0.4561", "0.0547", "0.0467")),
            roi_not_measured=compact("chưa có dữ liệu để điền tỷ lệ ROI") in full_compact
                or compact("không công bố tỷ lệ ROI") in full_compact,
        )
    if (ROOT / "docs/TICH_HOP_C3_NHAT_HUY.md").exists() and name in ("BAO_CAO_DO_AN_GOEMOTIONS_IEEE", "BAO_CAO_TIEN_DO_2"):
        verification = json.loads((REPORTS / "verification_project.json").read_text(encoding="utf-8"))
        tests = verification.get("post_merge_unit_tests", {})
        checks.update(
            huy_delivery_acknowledged="3acdfc6" in full and "app_distilbert_huy.py" in full,
            two_c3_studies_separate=all(value in full_compact for value in ("0.4061±0.0045", "0.4064±0.0060")),
            no_six_seed_pooling="sáu seed" in source and "validation" in source and "test" in source,
            post_merge_unit_tests_snapshot=f"{tests.get('passed')}/{tests.get('tests_run')}" in full_compact,
            old_unit_tests_snapshot_kept="69/69" in full_compact,
        )
    targets = {"title": 0}
    keywords = (("comparison", "BẢNG III."), ("cases", "eczj48j"), ("references", "TÀI LIỆU THAM KHẢO")) if name.startswith("BAI_BAO") else (
        ("comparison", "Bảng 5-2a."), ("cases", "eczj48j"), ("demo", "Hình 5.4."), ("references", "TÀI LIỆU THAM KHẢO"))
    if name.startswith("BAO_CAO_TIEN"):
        keywords = (("results", "A_standard"), ("roles", "Phân công" if name.endswith("1") else "Vai trò"))
    if name in ("BAO_CAO_DO_AN_GOEMOTIONS_IEEE", "BAO_CAO_TIEN_DO_2"):
        keywords += (("huy_integration", "3acdfc6"),)
    if name in ("BAO_CAO_DO_AN_GOEMOTIONS_IEEE", "BAI_BAO_GOEMOTIONS_IEEE"):
        keywords += (("hyperparameter_basis", "chưa thực hiện tìm kiếm siêu tham số có hệ thống"),
                     ("threshold_tradeoff", "test locked F1"),
                     ("applications_roi", "MTTR"))
    for label, keyword in keywords:
        candidates = [index for index, text in enumerate(page_texts) if compact(keyword) in compact(text)]
        if candidates:
            # Bản sáu chương có danh mục bảng chứa lại caption/ID ở đầu.
            use_last = label in ("demo", "references", "roles", "huy_integration") or name.startswith("BAO_CAO_DO_AN")
            targets[label] = candidates[-1] if use_last else candidates[0]
    previews = [preview(pdf, index, label, name) for label, index in targets.items()]
    pdf.close()
    return {"name": name, "source": source_name, "pdf_pages": len(page_texts), "docx_tables": len(word.tables),
            "docx_images": len(word.inline_shapes), "section_columns": cols,
            "normal_font": {"name": word.styles["Normal"].font.name, "size_pt": word.styles["Normal"].font.size.pt},
            "page_inches": [round(value / 72, 3) for value in sizes[0]],
            "citation_first_use_order": citation_order(source), "checks": checks,
            "checks_passed": all(checks.values()),
            "artifacts": [{"path": str(path.relative_to(ROOT)), "bytes": path.stat().st_size, "sha256": digest(path)}
                          for path in (source_path, docx_path, pdf_path)],
            "previews": previews}


def main():
    PREVIEWS.mkdir(parents=True, exist_ok=True)
    summary_path = REPORTS / "project_results/summary.json"
    summary = json.loads(summary_path.read_text(encoding="utf-8"))
    reports = [inspect(name, source, refs) for name, source, refs in zip(NAMES, SOURCES, (26, 20, 8, 5))]
    evidence = {
        "checked_at_utc": datetime.now(timezone.utc).isoformat(),
        "scope": "Actual DOCX structure, actual PDF text/pages, source citation order, targeted rendered page previews.",
        "summary": {"complete": summary["complete"], "records": len(summary["records"]),
                    "averages": len(summary["averages"]), "sha256": digest(summary_path)},
        "post_merge_verification_sha256": digest(REPORTS / "verification_project.json"),
        "reports": reports, "automatic_checks_passed": all(report["checks_passed"] for report in reports),
        "visual_review": {"status": "pending_manual_review", "note": "Open generated previews before marking visually reviewed."},
        "limits": ["Administrative details and contribution percentages remain blank for the group to confirm.",
                   "Paper uses the IEEE conference layout; this is a course manuscript, not a published IEEE paper.",
                   "Actual experiments were completed on Oct. 8; document rendering date may be later.",
                   "Three seed sample std is descriptive; F1 does not establish industrial ROI.",
                   "Timing includes Windows sleep interruptions; do not rank throughput from elapsed times.",
                   "Huy's three-seed C3 validation study is separate from the main C3 study; never pool six seeds or substitute validation for test.",
                   "No model, GPU, new inference, training or Word export performed by this verification script."],
    }
    revision_path = REPORTS / "execution/scientific_revision_before_10_10_2026.json"
    if revision_path.exists():
        revision = json.loads(revision_path.read_text(encoding="utf-8"))
        unchanged = {path: digest(ROOT / path) == expected
                     for path, expected in revision["sha256_before"].items()}
        evidence["scientific_content_revision"] = {
            "client_date": "2026-10-10", "timezone": "Asia/Bangkok",
            "scope": "Applications/ROI, validation/test tradeoffs and hyperparameter evidence in both final reports",
            "experiment_and_source_unchanged": all(unchanged.values()), "checks": unchanged,
            "before_record": str(revision_path.relative_to(ROOT)),
        }
        evidence["automatic_checks_passed"] = evidence["automatic_checks_passed"] and all(unchanged.values())
    target = REPORTS / "execution/final_report_verification.json"
    target.write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"evidence": str(target), "automatic_checks_passed": evidence["automatic_checks_passed"],
                      "pages": {report["name"]: report["pdf_pages"] for report in reports}}, ensure_ascii=False))
    failures = {report["name"]: [name for name, passed in report["checks"].items() if not passed]
                for report in reports if not report["checks_passed"]}
    if failures:
        raise SystemExit(json.dumps(failures, ensure_ascii=False))
    if not evidence["automatic_checks_passed"]:
        raise SystemExit("Experiment/protocol/source hashes changed during documentation revision")


if __name__ == "__main__":
    main()
